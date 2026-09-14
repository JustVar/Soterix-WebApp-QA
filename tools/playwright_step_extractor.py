#!/usr/bin/env python3
"""Extract human-readable test steps from Playwright test code.

Parses test files via AST, translates common Playwright/page-object patterns
into QA-friendly steps, and outputs them keyed by QA Touch case key.

Usage:
    python tools/playwright_step_extractor.py analyze     # Show what would be generated
    python tools/playwright_step_extractor.py generate    # Write steps JSON for bulk update
"""
import ast
import io
import json
import os
import re
import sys

if sys.stdout.encoding != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAPPING_FILE = os.path.join(ROOT, "config", "qatouch_case_mapping.json")
TESTS_DIR = os.path.join(ROOT, "tests", "ui")
OUTPUT_FILE = os.path.join(ROOT, "config", "playwright_generated_steps.json")

# ---------------------------------------------------------------------------
# Module context: maps test file paths to human-readable module names
# ---------------------------------------------------------------------------
FILE_MODULE_MAP = {
    "test_license_navigation.py": "Licenses",
    "test_license_search.py": "Licenses",
    "test_license_assignment.py": "Licenses",
    "test_license_permissions.py": "Licenses",
    "test_logout.py": "Logout",
    "test_login_security.py": "Login Security",
    "test_login_valid.py": "Login",
    "test_login_invalid.py": "Login",
    "test_notifications.py": "Notifications",
    "test_notifications_deeper.py": "Notifications",
    "test_maps.py": "Maps",
    "test_maps_deeper.py": "Maps",
    "test_admin_general.py": "Admin - General",
    "test_admin_rules.py": "Admin - Rules",
    "test_admin_sites.py": "Admin - Sites",
    "test_admin_users.py": "Admin - Users",
    "test_admin_event_wall.py": "Admin - Event Wall",
    "test_admin_audit_trail.py": "Admin - Audit Trail",
    "test_events_filters.py": "Events",
    "test_events_playback.py": "Events",
    "test_events_search.py": "Events",
    "test_search_filters.py": "Search",
    "test_search_drawing.py": "Search",
    "test_search_deeper.py": "Search",
    "test_search_categories.py": "Search",
    "test_incamera_search.py": "Search - In Camera",
    "test_views_basic.py": "Views",
    "test_views_bookmarks.py": "Views - Bookmarks",
    "test_views_drag_drop.py": "Views - Drag & Drop",
    "test_views_favorites.py": "Views - Favorites",
    "test_views_fullscreen.py": "Views - Fullscreen",
    "test_views_playback.py": "Views - Playback",
    "test_views_share.py": "Views - Share",
    "test_views_subtabs.py": "Views - Subtabs",
    "test_camera_sequences.py": "Camera Sequences",
    "test_media_navigation.py": "Media",
    "test_media_operations.py": "Media",
    "test_media_deeper.py": "Media",
    "test_chats.py": "Chats",
    "test_chats_deeper.py": "Chats",
    "test_insights_navigation.py": "Insights",
    "test_insights_widgets.py": "Insights",
    "test_insights_dashboard.py": "Insights",
}

# ---------------------------------------------------------------------------
# Pattern-based step extraction from AST nodes
# ---------------------------------------------------------------------------

def _unparse_safely(node):
    """Get source representation of an AST node."""
    try:
        return ast.unparse(node)
    except Exception:
        return "..."


def _extract_string_arg(call_node, idx=0):
    """Extract a string argument from a function call."""
    if idx < len(call_node.args):
        arg = call_node.args[idx]
        if isinstance(arg, ast.Constant) and isinstance(arg.value, str):
            return arg.value
    return None


def _get_attr_chain(node):
    """Get dotted attribute chain like 'self.page.locator'."""
    parts = []
    while isinstance(node, ast.Attribute):
        parts.append(node.attr)
        node = node.value
    if isinstance(node, ast.Name):
        parts.append(node.id)
    return ".".join(reversed(parts))


def _locator_description(call_node):
    """Try to describe what a locator targets from its selector string."""
    selector = _extract_string_arg(call_node)
    if not selector:
        return None

    # text= patterns
    texts = re.findall(r"text=([^,\"]+)", selector)
    if texts:
        return " / ".join(t.strip() for t in texts[:3])

    # has-text patterns
    has_texts = re.findall(r":has-text\(['\"]([^'\"]+)['\"]\)", selector)
    if has_texts:
        return " / ".join(has_texts[:3])

    # input/button with name
    name = re.search(r'\[name=["\']([^"\']+)', selector)
    if name:
        return f"{name.group(1)} field"

    # href patterns
    href = re.search(r'\[href\*=["\']([^"\']+)', selector)
    if href:
        return f"link to {href.group(1)}"

    # role patterns
    role = re.search(r'\[role=["\']([^"\']+)', selector)
    if role:
        return role.group(1)

    # img with alt or src
    alt = re.search(r'\[alt\*=["\']([^"\']+)', selector)
    if alt:
        return f"{alt.group(1)} icon"
    src = re.search(r'\[src\*=["\']([^"\']+)', selector)
    if src:
        return f"{src.group(1)} icon"

    # CSS class hints
    cls = re.search(r"\.([\w-]+)", selector)
    if cls:
        return cls.group(1).replace("-", " ").replace("_", " ")

    # p:text-is('...')
    text_is = re.search(r"text-is\(['\"]([^'\"]+)['\"]\)", selector)
    if text_is:
        return text_is.group(1)

    return selector[:50] if len(selector) > 50 else selector


def _method_to_action(method_name):
    """Convert a page object method name to a human-readable action."""
    # Convert camelCase/snake_case to words
    words = re.sub(r"([A-Z])", r" \1", method_name)
    words = words.replace("_", " ").strip().lower()
    # Capitalize first word
    if words:
        words = words[0].upper() + words[1:]
    return words


class StepExtractor(ast.NodeVisitor):
    """Extracts human-readable steps from a test function's AST."""

    def __init__(self, helper_bodies=None, module_name=""):
        self.steps = []
        self.helper_bodies = helper_bodies or {}
        self.module_name = module_name
        self._last_comment = None
        self._local_vars = {}  # var_name -> AST node (for locator assignments)

    def _add(self, action, expected=""):
        self.steps.append({"action": action, "expected": expected})

    def _process_call(self, node):
        """Process a Call node and generate a step."""
        chain = _get_attr_chain(node.func) if isinstance(node.func, ast.Attribute) else ""

        # --- Page navigation ---
        if chain.endswith(".goto"):
            url = _extract_string_arg(node)
            if url:
                self._add(f"Navigate to {url}")
            else:
                self._add("Navigate to the target URL")
            return True

        if chain.endswith(".reload"):
            self._add("Refresh the page")
            return True

        if chain.endswith(".go_back"):
            self._add("Click the browser Back button")
            return True

        # --- Skip waits, they're not user-visible steps ---
        if chain.endswith((".wait_for_timeout", ".wait_for_url",
                          ".wait_for_load_state", ".wait_for")):
            return True

        # --- Locator interactions ---
        if chain.endswith(".click"):
            # Walk up to find what's being clicked
            parent = node.func.value if isinstance(node.func, ast.Attribute) else None
            desc = self._describe_element(parent)
            if desc == "the element":
                # Try to get context from a preceding comment
                desc = self._last_comment or "the element"
            self._add(f"Click on {desc}")
            return True

        if chain.endswith(".hover"):
            parent = node.func.value if isinstance(node.func, ast.Attribute) else None
            desc = self._describe_element(parent)
            self._add(f"Hover over {desc}")
            return True

        if chain.endswith(".fill"):
            val = _extract_string_arg(node)
            parent = node.func.value if isinstance(node.func, ast.Attribute) else None
            desc = self._describe_element(parent)
            if val:
                self._add(f"Enter '{val}' in {desc}")
            else:
                self._add(f"Enter text in {desc}")
            return True

        if chain.endswith(".press"):
            key = _extract_string_arg(node)
            self._add(f"Press {key} key")
            return True

        if chain.endswith(".type"):
            val = _extract_string_arg(node)
            parent = node.func.value if isinstance(node.func, ast.Attribute) else None
            desc = self._describe_element(parent)
            self._add(f"Type '{val}' in {desc}")
            return True

        # --- expect() calls ---
        if chain.endswith((".to_be_visible", ".to_be_hidden", ".to_have_value",
                          ".to_have_text", ".to_have_count", ".to_contain_text")):
            # This is a verification - handled in visit_Expr via assert/expect
            return False

        # --- Page object method calls ---
        if isinstance(node.func, ast.Attribute):
            obj_chain = _get_attr_chain(node.func.value) if isinstance(node.func.value, (ast.Attribute, ast.Name)) else ""
            method = node.func.attr

            # Skip internal/private methods that aren't user actions
            if method.startswith("_") and method not in self.helper_bodies:
                # Check if it's a helper we can inline
                if method in self.helper_bodies:
                    return self._inline_helper(method)
                return False

            # Known page object navigation methods
            if method == "navigate":
                if self.module_name:
                    self._add(f"Navigate to the {self.module_name} page")
                else:
                    self._add("Navigate to the module page")
                return True

            if method.startswith("navigate_to_"):
                target = method.replace("navigate_to_", "").replace("_", " ").title()
                self._add(f"Navigate to {target}")
                return True

            if method.startswith("switch_to_"):
                target = method.replace("switch_to_", "").replace("_", " ").title()
                self._add(f"Switch to {target}")
                return True

            if method == "login":
                self._add("Enter valid credentials and log in")
                return True

            if method == "logout":
                self._add("Open profile menu and click Sign Out")
                return True

            if method.startswith("open_"):
                target = method.replace("open_", "").replace("_", " ")
                self._add(f"Open {target}")
                return True

            if method.startswith("close_"):
                target = method.replace("close_", "").replace("_", " ")
                self._add(f"Close {target}")
                return True

            if method.startswith("click_"):
                target = method.replace("click_", "").replace("_", " ")
                self._add(f"Click {target}")
                return True

            if method == "search":
                val = _extract_string_arg(node)
                if val:
                    self._add(f"Search for '{val}'")
                else:
                    self._add("Perform a search")
                return True

            if method == "clear_search":
                self._add("Clear the search field")
                return True

            if method.startswith("select_"):
                target = method.replace("select_", "").replace("_", " ")
                val = _extract_string_arg(node)
                if val:
                    self._add(f"Select {target}: '{val}'")
                else:
                    self._add(f"Select {target}")
                return True

            if method.startswith("expand_"):
                target = method.replace("expand_", "").replace("_", " ")
                self._add(f"Expand {target}")
                return True

            if method.startswith("collapse_"):
                target = method.replace("collapse_", "").replace("_", " ")
                self._add(f"Collapse {target}")
                return True

            if method.startswith("toggle_"):
                target = method.replace("toggle_", "").replace("_", " ")
                self._add(f"Toggle {target}")
                return True

            if method.startswith("search_"):
                target = method.replace("search_", "").replace("_", " ")
                val = _extract_string_arg(node)
                if val:
                    self._add(f"Search {target} for '{val}'")
                else:
                    self._add(f"Search {target}")
                return True

            if method.startswith("edit_"):
                target = method.replace("edit_", "").replace("_", " ")
                self._add(f"Edit {target}")
                return True

            if method.startswith("delete_"):
                target = method.replace("delete_", "").replace("_", " ")
                self._add(f"Delete {target}")
                return True

            if method.startswith("suspend_"):
                target = method.replace("suspend_", "").replace("_", " ")
                self._add(f"Suspend {target}")
                return True

            if method.startswith("cancel_"):
                target = method.replace("cancel_", "").replace("_", " ")
                self._add(f"Cancel {target}")
                return True

            if method.startswith("is_") or method.startswith("get_") or method.startswith("find_"):
                # Read-only checks, not user actions
                return False

            # Generic page object call - convert method name
            if obj_chain.startswith("self.") and not obj_chain.endswith(".page"):
                action = _method_to_action(method)
                if action and method not in ("count", "first", "last", "nth"):
                    self._add(action)
                    return True

        return False

    def _describe_element(self, node):
        """Describe what element is being interacted with."""
        if node is None:
            return "the element"

        # If it's a locator call, describe the selector
        if isinstance(node, ast.Call):
            chain = _get_attr_chain(node.func) if isinstance(node.func, ast.Attribute) else ""
            if "locator" in chain:
                desc = _locator_description(node)
                if desc:
                    return desc

            # .nth(), .first, .last
            if chain.endswith((".nth", ".first", ".last")):
                inner = node.func.value if isinstance(node.func, ast.Attribute) else None
                inner_desc = self._describe_element(inner)
                if inner_desc != "the element":
                    return inner_desc

            # Chained calls like locator(...).first
            if isinstance(node.func, ast.Attribute):
                inner_desc = self._describe_element(node.func.value)
                if inner_desc != "the element":
                    return inner_desc

        # Named attribute like self.license.search_input
        if isinstance(node, ast.Attribute):
            attr = node.attr
            # Skip generic attributes
            if attr in ("first", "last", "value"):
                if isinstance(node.value, (ast.Attribute, ast.Call)):
                    return self._describe_element(node.value)
            # Convert snake_case to readable
            return attr.replace("_", " ")

        # Local variable name - resolve or use name as description
        if isinstance(node, ast.Name):
            name = node.id
            if name in self._local_vars:
                resolved = self._describe_element(self._local_vars[name])
                if resolved != "the element":
                    return resolved
            # Use variable name as fallback
            readable = name.replace("_", " ")
            if readable not in ("self", "page"):
                return readable

        return "the element"

    def _inline_helper(self, method_name):
        """Inline a helper method's steps."""
        if method_name in self.helper_bodies:
            for stmt in self.helper_bodies[method_name]:
                self.visit(stmt)
            return True
        return False

    def visit_Assign(self, node):
        """Track local variable assignments for locator descriptions."""
        # Store variable assignments so we can resolve names later
        for target in node.targets:
            if isinstance(target, ast.Name):
                self._local_vars[target.id] = node.value
        # Don't generate steps for assignments

    def visit_Expr(self, node):
        """Handle expression statements (most test lines)."""
        if isinstance(node.value, ast.Call):
            call = node.value
            chain = _get_attr_chain(call.func) if isinstance(call.func, ast.Attribute) else ""

            # expect(...).to_be_visible() pattern
            if ".to_be_visible" in chain or ".to_be_hidden" in chain:
                inner = call.func.value  # the expect() call
                if isinstance(inner, ast.Call):
                    inner_chain = _get_attr_chain(inner.func) if isinstance(inner.func, ast.Attribute) else ""
                    if "expect" in inner_chain or "expect" in (_get_attr_chain(inner.func) if isinstance(inner.func, ast.Attribute) else ""):
                        if inner.args:
                            desc = self._describe_element(inner.args[0])
                            if ".to_be_visible" in chain:
                                self._add(f"Verify {desc} is visible", f"{desc} is displayed")
                            else:
                                self._add(f"Verify {desc} is hidden", f"{desc} is no longer visible")
                            return

            if ".to_have_value" in chain:
                inner = call.func.value
                if isinstance(inner, ast.Call) and inner.args:
                    desc = self._describe_element(inner.args[0])
                    val = _extract_string_arg(call)
                    if val:
                        self._add(f"Verify {desc} has value '{val}'", f"Value is '{val}'")
                    return

            if ".to_have_text" in chain or ".to_contain_text" in chain:
                inner = call.func.value
                if isinstance(inner, ast.Call) and inner.args:
                    desc = self._describe_element(inner.args[0])
                    val = _extract_string_arg(call)
                    if val:
                        self._add(f"Verify {desc} contains '{val}'", f"Text '{val}' is shown")
                    return

            # Regular call
            if self._process_call(call):
                return

            # Check for helper method calls like self._navigate()
            if isinstance(call.func, ast.Attribute):
                method = call.func.attr
                if method in self.helper_bodies:
                    self._inline_helper(method)
                    return

        self.generic_visit(node)

    def visit_Assert(self, node):
        """Handle assert statements."""
        msg = None
        if node.msg and isinstance(node.msg, ast.Constant):
            msg = node.msg.value
        elif node.msg and isinstance(node.msg, ast.JoinedStr):
            # f-string - just use the static parts
            parts = []
            for v in node.msg.values:
                if isinstance(v, ast.Constant):
                    parts.append(str(v.value))
            msg = "".join(parts).strip()

        if msg:
            # Clean up assertion messages to be verification steps
            action = msg
            if action.lower().startswith("expected"):
                action = "Verify: " + action
            elif not action.lower().startswith("verify"):
                action = "Verify " + action[0].lower() + action[1:]
            self._add(action, msg)
        else:
            # Try to describe the assertion from the test
            test_src = _unparse_safely(node.test)
            if "count()" in test_src:
                self._add("Verify elements are present", "Expected elements found")
            elif "is_login_page" in test_src:
                self._add("Verify login page is displayed", "Login page is shown")
            elif "is_logged_in" in test_src:
                self._add("Verify user is logged in", "User is authenticated")

    def visit_With(self, node):
        """Handle 'with' blocks like expect_download."""
        for item in node.items:
            if isinstance(item.context_expr, ast.Call):
                chain = _get_attr_chain(item.context_expr.func) if isinstance(item.context_expr.func, ast.Attribute) else ""
                if "expect_download" in chain:
                    self._add("Initiate file download")
                elif "expect_popup" in chain:
                    self._add("Wait for popup window")

        # Visit body
        for stmt in node.body:
            self.visit(stmt)

    def visit_For(self, node):
        """Handle for loops - describe the iteration."""
        # Check if iterating over known values
        if isinstance(node.iter, ast.List):
            items = [_unparse_safely(e) for e in node.iter.elts[:3]]
            if items:
                self._add(f"For each item ({', '.join(items)}...):")
        # Visit body once (don't repeat steps)
        for stmt in node.body:
            self.visit(stmt)

    def visit_If(self, node):
        """Handle if blocks - visit body, skip else."""
        for stmt in node.body:
            self.visit(stmt)

    def visit_Try(self, node):
        """Handle try blocks."""
        for stmt in node.body:
            self.visit(stmt)


# ---------------------------------------------------------------------------
# File parsing
# ---------------------------------------------------------------------------

def parse_test_file(filepath):
    """Parse a test file and extract steps for each test function.

    Returns: dict of {nodeid_suffix: {"docstring": str, "steps": list}}
    """
    with open(filepath, encoding="utf-8") as f:
        source = f.read()

    try:
        tree = ast.parse(source)
    except SyntaxError:
        return {}

    results = {}
    rel_path = os.path.relpath(filepath, ROOT).replace("\\", "/")
    filename = os.path.basename(filepath)

    # Collect classes and their methods
    for node in ast.walk(tree):
        if isinstance(node, ast.ClassDef):
            class_name = node.name
            helpers = {}

            # First pass: collect helper methods
            for item in node.body:
                if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    if item.name.startswith("_") and item.name != "__init__":
                        helpers[item.name] = item.body

            # Second pass: process test methods
            for item in node.body:
                if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    if item.name.startswith("test_"):
                        docstring = ast.get_docstring(item) or ""
                        nodeid = f"{rel_path}::{class_name}::{item.name}"
                        steps = _extract_steps(item, helpers, docstring, filename)
                        results[nodeid] = {
                            "docstring": docstring,
                            "steps": steps,
                        }

        # Top-level test functions (not in a class)
        if isinstance(node, ast.Module):
            for item in node.body:
                if isinstance(item, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    if item.name.startswith("test_"):
                        docstring = ast.get_docstring(item) or ""
                        nodeid = f"{rel_path}::{item.name}"
                        steps = _extract_steps(item, {}, docstring, filename)
                        results[nodeid] = {
                            "docstring": docstring,
                            "steps": steps,
                        }

    return results


def _extract_steps(func_node, helpers, docstring, filename):
    """Extract steps from a test function AST node."""
    module_name = FILE_MODULE_MAP.get(filename, "")

    steps = []

    # Step 1: Always start with "Log in to the application" since all tests
    # use authenticated_page fixture
    steps.append({
        "action": "Log in to the Soterix application with valid credentials",
        "expected": "Dashboard loads successfully",
    })

    # Step 2: Extract steps from code
    extractor = StepExtractor(helper_bodies=helpers, module_name=module_name)
    for stmt in func_node.body:
        # Skip docstring (first Expr with Constant string)
        if isinstance(stmt, ast.Expr) and isinstance(stmt.value, ast.Constant) and isinstance(stmt.value.value, str):
            continue
        # Reset last comment tracker
        extractor._last_comment = None
        extractor.visit(stmt)

    # Add extracted steps, deduplicating consecutive identical steps
    prev = None
    for s in extractor.steps:
        action = s["action"]
        # Clean up assertion messages: strip trailing format placeholders
        action = re.sub(r",?\s*got:?\s*$", "", action)
        action = re.sub(r"\s*--\s*$", "", action)
        s["action"] = action.strip()
        if s.get("expected"):
            s["expected"] = re.sub(r",?\s*got:?\s*$", "", s["expected"]).strip()

        if s["action"] != prev and s["action"]:
            steps.append(s)
            prev = s["action"]

    # If we only have the login step, add a generic step from docstring
    if len(steps) <= 1 and docstring:
        steps.append({
            "action": docstring,
            "expected": "Test passes successfully",
        })

    # Number the steps
    for i, s in enumerate(steps):
        s["n"] = i + 1

    return steps


# ---------------------------------------------------------------------------
# Title-based step generation for cases with no Playwright test
# ---------------------------------------------------------------------------

MODULE_NAV_STEPS = {
    "Add New Rule": [
        {"action": "Navigate to Admin from the left sidebar", "expected": "Admin section loads"},
        {"action": "Click on General tab", "expected": "General section is displayed"},
        {"action": "Click on Rules sub-tab", "expected": "Rules list is displayed"},
    ],
    "Camera Search Filters": [
        {"action": "Navigate to Search from the left sidebar", "expected": "Search page loads"},
        {"action": "Select a camera with search capability", "expected": "Camera search view loads"},
    ],
    "Camera Sequence Options": [
        {"action": "Navigate to Camera Sequences from the left sidebar", "expected": "Sequences page loads"},
    ],
    "Camera Sequence Save": [
        {"action": "Navigate to Camera Sequences from the left sidebar", "expected": "Sequences page loads"},
    ],
    "Camera Sequence Timing": [
        {"action": "Navigate to Camera Sequences from the left sidebar", "expected": "Sequences page loads"},
    ],
    "Events Filter": [
        {"action": "Navigate to Events from the left sidebar", "expected": "Events page loads"},
    ],
    "General": [
        {"action": "Navigate to Admin from the left sidebar", "expected": "Admin section loads"},
        {"action": "Click on General tab", "expected": "General section is displayed"},
    ],
    "Logout": [
        # No extra nav needed, user is already logged in
    ],
    "Maps": [
        {"action": "Navigate to Maps from the left sidebar", "expected": "Maps page loads"},
        {"action": "Select a site from the dropdown", "expected": "Site maps are displayed"},
    ],
    "Notification Panel": [
        # Notifications are in the header, no extra nav needed
    ],
    "Search in Camera": [
        {"action": "Navigate to Views from the left sidebar", "expected": "Views page loads"},
        {"action": "Select a camera and open in-camera search", "expected": "In-camera search panel opens"},
    ],
    "Share Live View - Camera": [
        {"action": "Navigate to Views from the left sidebar", "expected": "Views page loads"},
        {"action": "Select a camera tile", "expected": "Camera stream is displayed"},
    ],
    "Sites": [
        {"action": "Navigate to Admin from the left sidebar", "expected": "Admin section loads"},
        {"action": "Click on Sites tab", "expected": "Sites list is displayed"},
        {"action": "Click on a site to open site detail", "expected": "Site detail page loads"},
    ],
}


def _generate_steps_from_title(title, module):
    """Generate human-readable steps from just the test case title and module."""
    steps = []

    # Step 1: Login
    steps.append({
        "action": "Log in to the Soterix application with valid credentials",
        "expected": "Dashboard loads successfully",
    })

    # Step 2+: Module-specific navigation
    nav_steps = MODULE_NAV_STEPS.get(module, [])
    for ns in nav_steps:
        steps.append(dict(ns))

    # Parse the title to derive action and expected result
    title_clean = title.strip()

    # Extract the action from the title
    if title_clean.lower().startswith("verify "):
        # "Verify X does Y" -> action: "X", expected: "Y"
        action_part = title_clean[7:]  # strip "Verify "

        # Special patterns
        if " shows " in action_part.lower():
            parts = re.split(r"\s+shows\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": parts[0].strip(), "expected": ""})
            steps.append({
                "action": f"Verify {action_part}",
                "expected": parts[1].strip() if len(parts) > 1 else action_part,
            })
        elif " returns " in action_part.lower():
            parts = re.split(r"\s+returns\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": parts[0].strip(), "expected": ""})
            steps.append({
                "action": f"Verify {action_part}",
                "expected": parts[1].strip() if len(parts) > 1 else "Expected results are shown",
            })
        elif " opens " in action_part.lower():
            parts = re.split(r"\s+opens\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": f"Click on {parts[0].strip()}", "expected": ""})
            steps.append({
                "action": f"Verify {parts[1].strip() if len(parts) > 1 else 'element'} is displayed",
                "expected": parts[1].strip() if len(parts) > 1 else action_part,
            })
        elif " updates " in action_part.lower():
            parts = re.split(r"\s+updates\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": parts[0].strip(), "expected": ""})
            steps.append({
                "action": f"Verify {parts[1].strip() if len(parts) > 1 else 'results'} are updated",
                "expected": parts[1].strip() if len(parts) > 1 else action_part,
            })
        elif " closes " in action_part.lower() and " when " not in action_part.lower():
            parts = re.split(r"\s+closes\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": parts[0].strip(), "expected": ""})
            steps.append({
                "action": f"Verify {parts[1].strip() if len(parts) > 1 else 'element'} is closed",
                "expected": f"{parts[1].strip() if len(parts) > 1 else 'Element'} is no longer visible",
            })
        elif " lists " in action_part.lower():
            parts = re.split(r"\s+lists\s+", action_part, maxsplit=1, flags=re.IGNORECASE)
            steps.append({"action": f"Open {parts[0].strip()}", "expected": ""})
            steps.append({
                "action": f"Verify {parts[1].strip() if len(parts) > 1 else 'options'} are listed",
                "expected": parts[1].strip() if len(parts) > 1 else "Options are displayed",
            })
        else:
            # Generic: the whole title becomes the verification step
            steps.append({
                "action": title_clean,
                "expected": action_part,
            })
    else:
        # Non-verify title (like "Save Slip & Fall rule successfully")
        steps.append({
            "action": title_clean,
            "expected": "Action completes successfully",
        })

    # Number steps
    for i, s in enumerate(steps):
        s["n"] = i + 1

    return steps


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def load_case_mapping():
    """Load QA Touch case mapping and build reverse lookup."""
    with open(MAPPING_FILE, encoding="utf-8") as f:
        data = json.load(f)

    mappings = data.get("mappings", {})

    # Normalize nodeids: strip [chromium] suffix for matching
    key_to_nodeids = {}
    nodeid_to_key = {}
    for nodeid, info in mappings.items():
        case_key = info.get("case_key")
        if not case_key:
            continue
        # Strip parameterization suffix like [chromium]
        clean = re.sub(r"\[.*\]$", "", nodeid)
        nodeid_to_key[clean] = case_key
        key_to_nodeids.setdefault(case_key, []).append(clean)

    return nodeid_to_key, key_to_nodeids


def load_post_import_keys():
    """Load the 137 post-import case keys from QA Touch export."""
    import openpyxl
    QT = os.path.join(
        os.path.expanduser("~"), "Downloads",
        "TestCases_06-03-26-15_23_43.xlsx",
        "TestCases_06-03-26-15_23_43.xlsx"
    )
    wb = openpyxl.load_workbook(QT, read_only=True)
    ws = wb["CasesSheet"]
    keys = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        created = str(r[21]) if r[21] else ""
        if "2026-02-10" not in created:
            keys[r[0]] = {"title": r[5], "module": r[4]}
    wb.close()
    return keys


def main():
    cmd = sys.argv[1] if len(sys.argv) > 1 else "analyze"

    print("Loading case mapping...")
    nodeid_to_key, key_to_nodeids = load_case_mapping()
    print(f"  {len(nodeid_to_key)} unique nodeids mapped")

    print("Loading post-import case keys...")
    post_keys = load_post_import_keys()
    print(f"  {len(post_keys)} post-import cases")

    # Find which test files contain tests mapped to post-import cases
    print("\nParsing test files...")
    test_files = []
    for dirpath, _, filenames in os.walk(TESTS_DIR):
        for fn in filenames:
            if fn.startswith("test_") and fn.endswith(".py"):
                test_files.append(os.path.join(dirpath, fn))

    all_steps = {}  # case_key -> {title, steps}
    file_count = 0
    test_count = 0

    for tf in sorted(test_files):
        parsed = parse_test_file(tf)
        if not parsed:
            continue
        file_count += 1

        for nodeid, data in parsed.items():
            # Try to match this test to a post-import case key
            case_key = nodeid_to_key.get(nodeid)
            if case_key and case_key in post_keys:
                if case_key not in all_steps:  # First match wins
                    all_steps[case_key] = {
                        "title": post_keys[case_key]["title"],
                        "module": post_keys[case_key]["module"],
                        "steps": data["steps"],
                        "docstring": data["docstring"],
                        "nodeid": nodeid,
                    }
                    test_count += 1

    print(f"  Parsed {file_count} test files")
    print(f"  Generated steps for {test_count} / {len(post_keys)} post-import cases")

    # Show unmatched
    matched_keys = set(all_steps.keys())
    unmatched = {k: v for k, v in post_keys.items() if k not in matched_keys}
    mapped_but_no_parse = 0
    truly_unmapped = 0
    for k, _v in unmatched.items():
        if k in key_to_nodeids:
            mapped_but_no_parse += 1
        else:
            truly_unmapped += 1

    print(f"\n  Mapped but nodeid mismatch: {mapped_but_no_parse}")
    print(f"  No Playwright test (truly unmapped): {truly_unmapped}")

    # Generate title-based steps for the 35 unmapped cases
    for k, v in unmatched.items():
        steps = _generate_steps_from_title(v["title"], v["module"])
        all_steps[k] = {
            "title": v["title"],
            "module": v["module"],
            "steps": steps,
            "docstring": "",
            "nodeid": "(generated from title)",
        }

    print(f"  Total with steps (code + title-based): {len(all_steps)}")

    if cmd == "analyze":
        # Show sample output
        print("\n" + "=" * 70)
        print("  SAMPLE GENERATED STEPS")
        print("=" * 70)
        shown = 0
        for key, data in sorted(all_steps.items()):
            if shown >= 10:
                break
            print(f"\n[{key}] [{data['module']}] {data['title']}")
            print(f"  Docstring: {data['docstring'][:80]}")
            print(f"  Nodeid: {data['nodeid']}")
            for s in data["steps"]:
                exp = f" -> {s['expected']}" if s.get("expected") else ""
                print(f"  Step {s['n']}: {s['action']}{exp}")
            shown += 1

        # Show unmatched details
        if unmatched:
            print(f"\n{'=' * 70}")
            print(f"  UNMATCHED CASES ({len(unmatched)})")
            print("=" * 70)
            for k, v in sorted(unmatched.items(), key=lambda x: x[1]["module"]):
                status = "mapped (nodeid mismatch)" if k in key_to_nodeids else "NO test"
                print(f"  [{k}] [{v['module']}] {v['title'][:55]} -- {status}")

    elif cmd == "generate":
        # Write output
        output = {
            "source": "playwright_code",
            "generated_at": __import__("datetime").datetime.now().isoformat(),
            "cases": {}
        }
        for key, data in all_steps.items():
            output["cases"][key] = {
                "title": data["title"],
                "module": data["module"],
                "steps": data["steps"],
            }

        with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
            json.dump(output, f, indent=2, ensure_ascii=False)

        print(f"\nWrote {len(output['cases'])} cases to {OUTPUT_FILE}")

    return all_steps


if __name__ == "__main__":
    main()
