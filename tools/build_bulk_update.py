#!/usr/bin/env python3
"""Build QA Touch Bulk Update XLSX with Excel steps and duplicate title fixes."""

import io
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

if sys.stdout.encoding != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

import openpyxl

REPO = Path(__file__).resolve().parent.parent
QT_EXPORT = Path(r"C:\Users\JustinVarughese\Downloads\TestCases_06-03-26-15_23_43.xlsx\TestCases_06-03-26-15_23_43.xlsx")
EXCEL_IMPORT = Path(r"C:\Users\JustinVarughese\Downloads\Soterix_Testcases (1).xlsx")
OUTPUT = Path(r"C:\Users\JustinVarughese\Downloads\QATouch_BulkUpdate_Steps.xlsx")
PLAYWRIGHT_STEPS = REPO / "config" / "playwright_generated_steps.json"

# Standard column indices for original Excel
COLS = {"scenario": 0, "acceptance": 9, "steps": 10, "feature": 4, "sub_feature": 5}
COLS_OFFSET = {k: v + 1 for k, v in COLS.items()}

# Sheet-to-module mapping for disambiguating duplicates
SHEET_MODULE_MAP = {
    "Views": ["Audio", "Configuration and Settings"],
    "Events": ["Review Pinned Clips - Media"],
    "Search": ["Review Pinned Clips - Search"],
    "Admin-Rules": ["Entry-Exit", "Fire-Smoke Detection"],
    "Admin-Licenses": ["Licenses"],
    "Playback autotimeout": ["Playback Auto-Timeout"],
}

VERIFY_PREFIXES = [
    "Verify that ", "Verify ", "Validate that ", "Validate ",
    "Confirm that ", "Confirm ", "Ensure that ", "Ensure ",
    "Check that ", "Check ",
]


def norm(t):
    t = unicodedata.normalize("NFKD", str(t)).strip()
    t = re.sub(r"\s+", " ", t)
    t = t.replace("\u22ee", "...").replace("\u2026", "...")
    t = t.replace("\u2192", "-").replace("\u2190", "-")
    t = t.replace("\u2013", "-").replace("\u2014", "-")
    # Normalize smart quotes to straight quotes
    t = t.replace("\u201c", '"').replace("\u201d", '"')  # double
    t = t.replace("\u2018", "'").replace("\u2019", "'")  # single
    return t.lower().strip()


def parse_steps_text(steps_text, acceptance=""):
    if not steps_text or not steps_text.strip():
        return []
    lines = re.split(r"\n(?=Step\s+\d+\s*:)", steps_text.strip())
    steps = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        cleaned = re.sub(r"^Step\s+\d+\s*:\s*", "", line).strip()
        cleaned = re.sub(r"^Step\s+\d+\s*:\s*", "", cleaned).strip()
        if not cleaned:
            continue
        steps.append(cleaned)
    if not steps:
        for line in steps_text.strip().split("\n"):
            line = line.strip()
            if line:
                cleaned = re.sub(r"^\d+[.)]\s*", "", line).strip()
                if cleaned:
                    steps.append(cleaned)
    if not steps:
        return []

    result = []
    has_verify = any(_is_verify(s) for s in steps)
    for i, action in enumerate(steps, 1):
        expected = ""
        if _is_verify(action):
            expected = _derive_expected(action)
            if i == len(steps) and acceptance and acceptance.strip():
                ac = acceptance.strip().rstrip(".")
                if expected and ac.lower() != expected.lower():
                    expected = f"{expected}. {ac}"
                elif not expected:
                    expected = ac
        elif i == len(steps) and not has_verify and acceptance and acceptance.strip():
            expected = acceptance.strip().rstrip(".")
        result.append({"n": i, "action": action, "expected": expected})
    return result


def _is_verify(text):
    lower = text.lower()
    return any(lower.startswith(p.lower()) for p in VERIFY_PREFIXES)


def _derive_expected(text):
    for prefix in VERIFY_PREFIXES:
        if text.lower().startswith(prefix.lower()):
            derived = text[len(prefix):].strip()
            if derived:
                return derived[0].upper() + derived[1:]
    return ""


def load_excel_rows():
    wb = openpyxl.load_workbook(EXCEL_IMPORT, read_only=True, data_only=True)
    rows = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        cols = COLS
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header and header[0] and "test case id" in str(header[0]).lower():
            cols = COLS_OFFSET
        for row in ws.iter_rows(min_row=2, values_only=True):
            sc_idx = cols["scenario"]
            if sc_idx >= len(row) or not row[sc_idx]:
                continue
            scenario = str(row[sc_idx]).strip()
            if not scenario:
                continue

            def _get(key):
                idx = cols[key]
                return str(row[idx]).strip() if idx < len(row) and row[idx] else ""

            rows.append({
                "scenario": scenario,
                "acceptance": _get("acceptance"),
                "steps_text": _get("steps"),
                "feature": _get("feature"),
                "sub_feature": _get("sub_feature"),
                "sheet": sn,
            })
    wb.close()
    return rows


def load_qt_export():
    wb = openpyxl.load_workbook(QT_EXPORT, read_only=True)
    ws = wb["CasesSheet"]
    cases = []
    for r in list(ws.iter_rows(min_row=2, values_only=True)):
        cases.append({
            "key": r[0],
            "title": str(r[5]).strip() if r[5] else "",
            "module": str(r[4]) if r[4] else "",
        })
    wb.close()
    return cases


SKIP_RENAME_KEYS = {
    "aj6n7a",  # Accidental duplicate of qzdKve, created by scaffold tool - delete from QA Touch
}


def build_rename_map(qt_cases):
    title_counts = Counter(c["title"].lower().strip() for c in qt_cases)
    dupe_titles = {t for t, count in title_counts.items() if count > 1}
    rename_map = {}
    for dt in dupe_titles:
        cases = [c for c in qt_cases if c["title"].lower().strip() == dt]
        # Skip rename if all cases in the group are in the skip list or only one remains
        active = [c for c in cases if c["key"] not in SKIP_RENAME_KEYS]
        if len(active) <= 1:
            continue
        modules = [c["module"] for c in active]
        if len(set(modules)) > 1:
            for c in active:
                rename_map[c["key"]] = f"[{c['module']}] {c['title']}"
        else:
            for i, c in enumerate(active):
                rename_map[c["key"]] = f"{c['title']} ({i + 1})"
    return rename_map


def match_excel_to_qt(excel_rows, qt_cases):
    qt_index = {}
    for c in qt_cases:
        t = norm(c["title"])
        qt_index.setdefault(t, []).append(c)

    matched = {}
    matched_keys = set()
    unmatched = []

    for row in excel_rows:
        if not row["steps_text"]:
            continue
        scenario = row["scenario"]
        acceptance = row["acceptance"]
        ns = norm(scenario)
        ck = None

        # Tier 1: Composite
        if acceptance:
            composite = norm(f"{scenario} - {acceptance}")
            hits = [c for c in qt_index.get(composite, []) if c["key"] not in matched_keys]
            if hits:
                ck = hits[0]["key"]

        # Tier 2: Plain (unique or module-disambiguated)
        if not ck:
            hits = [c for c in qt_index.get(ns, []) if c["key"] not in matched_keys]
            if len(hits) == 1:
                ck = hits[0]["key"]
            elif len(hits) > 1:
                # Use sheet to pick the right module
                preferred = SHEET_MODULE_MAP.get(row["sheet"], [])
                for pm in preferred:
                    for h in hits:
                        if pm.lower() in h["module"].lower():
                            ck = h["key"]
                            break
                    if ck:
                        break
                if not ck and row["feature"]:
                    for h in hits:
                        if (row["feature"].lower() in h["module"].lower()
                                or h["module"].lower() in row["feature"].lower()):
                            ck = h["key"]
                            break
                if not ck:
                    ck = hits[0]["key"]

        # Tier 3: Prefix
        if not ck and len(ns) >= 15:
            for t, cases in qt_index.items():
                if t.startswith(ns) and len(ns) / max(len(t), 1) >= 0.6:
                    hits = [c for c in cases if c["key"] not in matched_keys]
                    if len(hits) == 1:
                        ck = hits[0]["key"]
                        break

        # Tier 4: QT title starts with scenario + " - "
        if not ck:
            for t, cases in qt_index.items():
                if t.startswith(ns + " - "):
                    hits = [c for c in cases if c["key"] not in matched_keys]
                    if len(hits) == 1:
                        ck = hits[0]["key"]
                        break

        # Tier 5: Arrow normalization
        if not ck:
            ns_alt = ns.replace("->", "-").replace(" -> ", " - ")
            if ns_alt != ns:
                hits = [c for c in qt_index.get(ns_alt, []) if c["key"] not in matched_keys]
                if len(hits) == 1:
                    ck = hits[0]["key"]
                if not ck:
                    for t, cases in qt_index.items():
                        if t.startswith(ns_alt):
                            hits = [c for c in cases if c["key"] not in matched_keys]
                            if len(hits) == 1:
                                ck = hits[0]["key"]
                                break

        if ck:
            matched[ck] = row
            matched_keys.add(ck)
        else:
            unmatched.append(row)

    return matched, unmatched


def main():
    print("Loading Excel import data...")
    excel_rows = load_excel_rows()
    print(f"  {len(excel_rows)} rows")

    print("Loading QA Touch export...")
    qt_cases = load_qt_export()
    print(f"  {len(qt_cases)} cases")

    print("Building rename map for duplicate titles...")
    rename_map = build_rename_map(qt_cases)
    print(f"  {len(rename_map)} titles to rename")
    for ck, new_title in rename_map.items():
        module = next((c["module"] for c in qt_cases if c["key"] == ck), "?")
        print(f"    [{ck}] [{module}] -> {new_title[:70]}")

    print("\nMatching Excel rows to QA Touch cases...")
    matched, unmatched = match_excel_to_qt(excel_rows, qt_cases)
    print(f"  Matched: {len(matched)}")
    print(f"  Unmatched: {len(unmatched)}")
    if unmatched:
        for r in unmatched:
            print(f"    [{r['sheet']}] {r['scenario'][:60]}")

    # Load Playwright-generated steps for cases not in Excel
    pw_steps = {}
    if PLAYWRIGHT_STEPS.exists():
        print("Loading Playwright-generated steps...")
        with open(PLAYWRIGHT_STEPS, "r", encoding="utf-8") as f:
            pw_data = json.load(f)
        pw_steps = pw_data.get("cases", {})
        print(f"  {len(pw_steps)} cases from Playwright code")

    print("\nBuilding bulk update XLSX...")
    wb_src = openpyxl.load_workbook(QT_EXPORT)
    ws_src = wb_src["CasesSheet"]
    src_rows = list(ws_src.iter_rows(values_only=False))

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "CasesSheet"

    # Header
    for ci, cell in enumerate(src_rows[0]):
        out_ws.cell(row=1, column=ci + 1, value=cell.value)

    updated_excel = 0
    updated_pw = 0
    renamed = 0
    for ri, row in enumerate(src_rows[1:], 2):
        case_key = row[0].value
        for ci, cell in enumerate(row):
            out_ws.cell(row=ri, column=ci + 1, value=cell.value)

        # Apply rename
        if case_key in rename_map:
            out_ws.cell(row=ri, column=6, value=rename_map[case_key])
            renamed += 1

        # Apply Excel steps (priority 1)
        if case_key in matched:
            excel_row = matched[case_key]
            steps = parse_steps_text(excel_row["steps_text"], excel_row.get("acceptance", ""))
            if steps:
                step_lines = [s["action"] for s in steps]
                exp_lines = [s.get("expected", "") for s in steps]
                out_ws.cell(row=ri, column=12, value="\n".join(step_lines))
                out_ws.cell(row=ri, column=13, value="\n".join(exp_lines))
                updated_excel += 1
        # Apply Playwright-generated steps (priority 2 - only if no Excel match)
        elif case_key in pw_steps:
            pw_case = pw_steps[case_key]
            steps = pw_case.get("steps", [])
            if steps:
                step_lines = [s["action"] for s in steps]
                exp_lines = [s.get("expected", "") for s in steps]
                out_ws.cell(row=ri, column=12, value="\n".join(step_lines))
                out_ws.cell(row=ri, column=13, value="\n".join(exp_lines))
                updated_pw += 1

    out_wb.save(OUTPUT)
    print(f"\nDone!")
    print(f"  Steps from Excel: {updated_excel}")
    print(f"  Steps from Playwright: {updated_pw}")
    print(f"  Steps total: {updated_excel + updated_pw}")
    print(f"  Titles renamed: {renamed}")
    print(f"  Total cases: {len(src_rows) - 1}")
    print(f"  Saved to: {OUTPUT}")

    out_wb.close()
    wb_src.close()


if __name__ == "__main__":
    main()
